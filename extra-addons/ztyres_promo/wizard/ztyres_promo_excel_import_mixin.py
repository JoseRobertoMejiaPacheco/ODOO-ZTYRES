# -*- coding: utf-8 -*-
"""
Base común de los asistentes de importación desde Excel.

Concentra tres cosas que estaban duplicadas o faltaban:

1. **Lectura tolerante de encabezados.** Antes el archivo tenía que traer
   exactamente `codigo` y `monto`, en minúsculas y sin acentos. Un
   archivo con "Código" o "CODIGO" o "Monto NC" fallaba con un mensaje
   que no decía qué encabezados sí traía. Ahora se normalizan acentos,
   mayúsculas y espacios, se aceptan sinónimos habituales, y el error
   lista las columnas encontradas.

2. **Plantilla descargable.** Cada asistente genera su propio .xlsx de
   ejemplo, con la hoja de datos primero (pandas lee la primera hoja) y
   una hoja de instrucciones después.

3. **Notificación de resultado** con el mismo formato en ambos.
"""

import base64
import io
import unicodedata

from odoo import _, fields, models
from odoo.exceptions import UserError


def normalize_header(value):
    """'Código del Producto ' -> 'codigo_del_producto'."""
    text = str(value or '').strip().lower()
    text = ''.join(
        character
        for character in unicodedata.normalize('NFKD', text)
        if not unicodedata.combining(character)
    )
    return '_'.join(text.split())


# Sinónimos aceptados para cada columna esperada.
COLUMN_ALIASES = {
    'codigo': (
        'codigo',
        'code',
        'clave',
        'sku',
        'default_code',
        'codigo_producto',
        'codigo_del_producto',
        'codigo_interno',
        'referencia',
        'referencia_interna',
    ),
    'monto': (
        'monto',
        'importe',
        'amount',
        'monto_nc',
        'monto_cupon',
        'monto_del_cupon',
        'monto_por_pieza',
        'valor',
    ),
    'desde': (
        'desde',
        'limite_inferior',
        'lower_limit',
        'minimo',
        'desde_monto',
        'desde_cantidad',
        'inicio',
    ),
    'hasta': (
        'hasta',
        'limite_superior',
        'upper_limit',
        'maximo',
        'hasta_monto',
        'hasta_cantidad',
        'fin',
    ),
    'porcentaje': (
        'porcentaje',
        'porciento',
        'descuento',
        'discount',
        'percent',
        'porcentaje_de_descuento',
        'porcentaje_descuento',
    ),
}


class ExcelImportMixin(models.AbstractModel):
    _name = 'ztyres_promo.excel_import_mixin'
    _description = 'Base para importaciones desde Excel'

    file = fields.Binary(string='Archivo Excel', required=True)
    file_name = fields.Char(string='Nombre del archivo')

    # ------------------------------------------------------------------
    # Lectura
    # ------------------------------------------------------------------
    def _read_excel(self, required_columns=()):
        """Lee la primera hoja y normaliza los encabezados."""
        self.ensure_one()
        if not self.file:
            raise UserError(_('Debe subir un archivo Excel.'))

        try:
            import pandas as pd
            data = pd.read_excel(io.BytesIO(base64.b64decode(self.file)))
        except UserError:
            raise
        except Exception as error:
            raise UserError(
                _('No fue posible leer el archivo Excel: %s') % error
            ) from error

        original_headers = list(data.columns)
        data.columns = [
            self._resolve_column(header) for header in data.columns
        ]

        missing = [
            column
            for column in required_columns
            if column not in data.columns
        ]
        if missing:
            raise UserError(_(
                'Al archivo le faltan estas columnas: %(faltan)s.\n\n'
                'Columnas encontradas: %(encontradas)s.\n\n'
                'Use el botón "Descargar plantilla" para ver el formato '
                'esperado.'
            ) % {
                'faltan': ', '.join(missing),
                'encontradas': ', '.join(
                    str(header) for header in original_headers
                ) or _('ninguna'),
            })
        return data

    def _resolve_column(self, header):
        normalized = normalize_header(header)
        for canonical, aliases in COLUMN_ALIASES.items():
            if normalized in aliases:
                return canonical
        return normalized

    # ------------------------------------------------------------------
    # Plantilla descargable
    # ------------------------------------------------------------------
    def _template_definition(self):
        """Debe devolver (nombre_archivo, columnas, filas, instrucciones)."""
        raise NotImplementedError

    def action_download_template(self):
        self.ensure_one()
        file_name, columns, rows, instructions = self._template_definition()
        content = self._build_template_file(
            columns,
            rows,
            instructions,
        )
        attachment = self.env['ir.attachment'].create({
            'name': file_name,
            'type': 'binary',
            'datas': base64.b64encode(content),
            'mimetype': (
                'application/vnd.openxmlformats-officedocument'
                '.spreadsheetml.sheet'
            ),
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'self',
        }

    def _build_template_file(self, columns, rows, instructions):
        """Genera el .xlsx. La hoja de datos va primero a propósito.

        `pandas.read_excel` lee la primera hoja del libro: si las
        instrucciones fueran la hoja 1, la importación leería el
        instructivo en lugar de los datos.
        """
        try:
            import xlsxwriter
        except ImportError:
            return self._build_template_file_openpyxl(
                columns,
                rows,
                instructions,
            )

        output = io.BytesIO()
        workbook = xlsxwriter.Workbook(output, {'in_memory': True})
        header_format = workbook.add_format({
            'bold': True,
            'bg_color': '#1F3864',
            'font_color': '#FFFFFF',
            'border': 1,
        })
        cell_format = workbook.add_format({'border': 1})
        title_format = workbook.add_format({'bold': True, 'font_size': 12})
        wrap_format = workbook.add_format({'text_wrap': True, 'valign': 'top'})

        sheet = workbook.add_worksheet('datos')
        for index, column in enumerate(columns):
            sheet.write(0, index, column, header_format)
            sheet.set_column(index, index, max(18, len(column) + 4))
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row):
                sheet.write(row_index, column_index, value, cell_format)
        sheet.freeze_panes(1, 0)

        guide = workbook.add_worksheet('instrucciones')
        guide.set_column(0, 0, 110)
        guide.write(0, 0, 'Cómo llenar esta plantilla', title_format)
        for index, line in enumerate(instructions, start=2):
            guide.write(index, 0, line, wrap_format)

        workbook.close()
        return output.getvalue()

    def _build_template_file_openpyxl(self, columns, rows, instructions):
        """Respaldo cuando xlsxwriter no está disponible."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'datos'
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F3864')
        for row in rows:
            sheet.append(list(row))
        for index, column in enumerate(columns, start=1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=index).column_letter
            ].width = max(18, len(column) + 4)
        sheet.freeze_panes = 'A2'

        guide = workbook.create_sheet('instrucciones')
        guide.column_dimensions['A'].width = 110
        guide.append(['Cómo llenar esta plantilla'])
        guide['A1'].font = Font(bold=True, size=12)
        for line in instructions:
            guide.append([line])
        for row in guide.iter_rows(min_col=1, max_col=1):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical='top')

        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()

    # ------------------------------------------------------------------
    # Resultado
    # ------------------------------------------------------------------
    def _notification(self, message, title=None, kind='success'):
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': title or _('Importación completada'),
                'message': message,
                'type': kind,
                'sticky': kind != 'success',
                'next': {'type': 'ir.actions.act_window_close'},
            },
        }
