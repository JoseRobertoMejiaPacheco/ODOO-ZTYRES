# -*- coding: utf-8 -*-
"""Banco de pruebas de los dos Excel del cotizador.

Genera lista de precios y pedido con datos ficticios, vuelve a abrir ambos
archivos y verifica contenido, distribución y formatos esenciales. Para ver
los archivos físicamente durante una ejecución de pruebas, definir:

    ZTYRES_XLSX_FIXTURES_DIR=/tmp/cotizador_xlsx
"""
import os
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from odoo.tests import tagged

from odoo.addons.ztyres_promo.models.ztyres_promo_config import (
    POLICY_QUANTITY,
    REWARD_PERCENTAGE,
    SCOPE_SPECIFIC_PRODUCTS,
    UNLIMITED,
)
from odoo.addons.ztyres_promo.tests.common import PromoTestCommon


@tagged('post_install', '-at_install')
class TestCotizadorXlsx(PromoTestCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.promo_excel = cls._promo(
            nombre='GOODYEAR AGOSTO 7.5%',
            promo_conditions=SCOPE_SPECIFIC_PRODUCTS,
            product_ids=[(6, 0, [cls.p_goodyear_16.product_tmpl_id.id])],
            promo_type=POLICY_QUANTITY,
            reward_type=REWARD_PERCENTAGE,
            policy_line_ids=[(0, 0, {
                'lower_limit': 12,
                'upper_limit': UNLIMITED,
                'discount': 7.5,
            })],
        )

    def _state(self):
        return {
            'partner_id': self.partner.id,
            'profile': {'volumen': 1, 'logistico': 2, 'financiero': 3},
            'promo_sim': {str(self.promo_excel.id): {'on': True, 'tier': 0}},
            'show_iva': True,
            'cart': {str(self.p_goodyear_16.id): 12},
            'search': '',
            'sf_filters': {},
        }

    def _generate_pair(self):
        Order = self.env['sale.order']
        files = {
            'lista_precios_datos_ficticios.xlsx': Order.download_pricelist_xlsx(self._state()),
            'pedido_datos_ficticios.xlsx': Order.download_order_xlsx(self._state()),
        }
        output_dir = os.environ.get('ZTYRES_XLSX_FIXTURES_DIR')
        if output_dir:
            destination = Path(output_dir)
            destination.mkdir(parents=True, exist_ok=True)
            for filename, content in files.items():
                (destination / filename).write_bytes(content)
        return files

    def test_genera_y_valida_los_dos_excel(self):
        files = self._generate_pair()
        self.assertEqual(set(files), {
            'lista_precios_datos_ficticios.xlsx',
            'pedido_datos_ficticios.xlsx',
        })
        for filename, content in files.items():
            workbook = load_workbook(BytesIO(content), data_only=False)
            sheet = workbook.active
            values = [
                cell.value
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            ]
            self.assertTrue(any(
                'Obtén 7.5% en la compra de 12 pzas o más.' in value
                for value in values
            ), filename)
            self.assertEqual(sheet.sheet_view.showGridLines, False)
            self.assertEqual(sheet.page_setup.orientation, 'landscape')
            self.assertTrue(sheet.auto_filter.ref)
            self.assertTrue(sheet.freeze_panes)
            self.assertTrue(sheet.print_title_rows)
            self.assertEqual(sheet['G4'].alignment.horizontal, 'right')
            self.assertEqual(sheet['G5'].alignment.horizontal, 'right')
            self.assertEqual(sheet['E3'].alignment.horizontal, 'right')
            self.assertEqual(sheet['I4'].alignment.horizontal, 'right')
            self.assertEqual(sheet['I5'].alignment.horizontal, 'right')
            self.assertTrue(sheet['I4'].alignment.wrap_text)
            self.assertTrue(sheet['I5'].alignment.wrap_text)
            last_letter = sheet.cell(1, sheet.max_column).column_letter
            merged = {str(cell_range) for cell_range in sheet.merged_cells.ranges}
            self.assertIn('G4:H4', merged)
            self.assertIn('G5:H5', merged)
            self.assertIn(f'E3:{last_letter}3', merged)
            self.assertIn(f'I4:{last_letter}4', merged)
            self.assertIn(f'I5:{last_letter}5', merged)

    def test_descripciones_iguales_al_estilo_de_la_pagina(self):
        Order = self.env['sale.order']
        promo = {
            'promo_type': 'quantity',
            'reward_type': 'percentage',
        }
        tier = {'min': 12, 'max': UNLIMITED, 'discount': 7.5, 'min_qty': 12}
        self.assertEqual(
            Order._cotizador_promo_detail(promo, tier),
            'Obtén 7.5% en la compra de 12 pzas o más.',
        )

    # ==================================================================
    # Precio base con Mayoreo, tarjeta de regalo fuera del precio
    # ==================================================================
    def test_precio_base_ya_trae_el_descuento_de_mayoreo(self):
        """El −10% de Mayoreo define el precio BASE, no un renglón final.

        Antes el precio publicado se mostraba tal cual y el −10% se
        restaba junto con los demás descuentos al final. El vendedor veía
        en la columna Precio un número que nunca se cobraba, y la promo
        se calculaba sobre esa base inflada.
        """
        Order = self.env['sale.order']
        con_mayoreo = {'price': 1000.0, 'mayoreo_discount': 10.0}
        sin_mayoreo = {'price': 1000.0, 'mayoreo_discount': 0.0}

        self.assertEqual(Order._cotizador_base_price(con_mayoreo), 900.0)
        self.assertEqual(Order._cotizador_base_price(sin_mayoreo), 1000.0)

    def test_promo_se_calcula_sobre_la_base_con_mayoreo(self):
        """7.5% sobre 900, no sobre 1000: 67.50 y no 75."""
        Order = self.env['sale.order']
        product_row = {
            'product_id': self.p_goodyear_16.id,
            'tmpl_id': self.p_goodyear_16.product_tmpl_id.id,
            'price': 900.0,          # ya viene de _cotizador_base_price
            'mayoreo_discount': 10.0,
        }
        promo = {
            'promo_type': 'quantity',
            'reward_type': 'percentage',
            'product_ids': [self.p_goodyear_16.id],
        }
        tier = {'min': 12, 'max': UNLIMITED, 'discount': 7.5}
        self.assertAlmostEqual(
            Order._cotizador_promo_discount(product_row, promo, tier),
            67.5,
            places=2,
        )

    def test_tarjeta_de_regalo_no_descuenta_el_precio(self):
        """La tarjeta es un valor que el cliente RECIBE.

        Descontarla del precio unitario hacía ver la llanta más barata de
        lo que se cobra — el precio con descuento potencial salía
        engañoso. Sale del precio y va a su propia columna.
        """
        Order = self.env['sale.order']
        product_row = {
            'product_id': self.p_goodyear_16.id,
            'tmpl_id': self.p_goodyear_16.product_tmpl_id.id,
            'price': 1000.0,
            'mayoreo_discount': 0.0,
        }
        promo = {
            'promo_type': 'quantity',
            'reward_type': 'gift_card',
            'product_ids': [self.p_goodyear_16.id],
            'coupons': [{
                'tmpl_id': self.p_goodyear_16.product_tmpl_id.id,
                'amount': 84.07,
            }],
        }
        tier = {'min': 1, 'max': UNLIMITED, 'discount': 0.0}

        self.assertEqual(
            Order._cotizador_promo_discount(product_row, promo, tier), 0.0)
        self.assertAlmostEqual(
            Order._cotizador_gift_card_amount(product_row, promo), 84.07,
            places=2)

    def test_cupon_por_pieza_si_sigue_descontando(self):
        """El cupón/apoyo fijo NO es tarjeta de regalo: sí baja el precio.

        Tiene tope de piezas y se refleja en la NC del pedido
        (docs/TARJETA_REGALO.md contrasta ambos). Solo se sacó del precio
        la tarjeta de regalo.
        """
        Order = self.env['sale.order']
        product_row = {
            'product_id': self.p_goodyear_16.id,
            'tmpl_id': self.p_goodyear_16.product_tmpl_id.id,
            'price': 1000.0,
        }
        promo = {
            'promo_type': 'coupons',
            'reward_type': 'percentage',
            'product_ids': [self.p_goodyear_16.id],
            'coupons': [{
                'tmpl_id': self.p_goodyear_16.product_tmpl_id.id,
                'amount': 50.0,
            }],
        }
        self.assertEqual(
            Order._cotizador_promo_discount(product_row, promo, None), 50.0)
        self.assertEqual(
            Order._cotizador_gift_card_amount(product_row, promo), 0.0)

    def test_los_dos_excel_traen_columna_de_tarjeta_de_regalo(self):
        """La columna existe en ambos archivos y la leyenda la explica."""
        for filename, content in self._generate_pair().items():
            sheet = load_workbook(BytesIO(content), data_only=False).active
            headers = [
                cell.value
                for row in sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str)
            ]
            self.assertIn('Tarjeta de regalo', headers, filename)
            self.assertIn('Precio', headers, filename)
            self.assertNotIn('Precio de lista', headers, filename)
            self.assertTrue(any(
                'NO descuenta el precio de las llantas' in value
                for value in headers
            ), filename)
